# Builds docs/innomate-section-sources.xlsx
# Sheets: ملخص · خريطة السكاشن · مراجعة المواقع · المنافسون · الكلمات المفتاحية · صفحات الموقع · مؤقت للتأكيد
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
content = json.load(open(os.path.join(HERE, 'content.json'), encoding='utf-8'))
refs_path = os.path.join(ROOT, 'research', 'references.json')
refs = json.load(open(refs_path, encoding='utf-8')) if os.path.exists(refs_path) else []

INK = '0A1A2B'; BLUE = '2B7BBD'; AMBER = 'F3A93A'; PAPER = 'F4F7FA'
FONT = 'Arial'
thin = Side(style='thin', color='DDE5EE')
wb = Workbook()
wb.remove(wb.active)

def sheet(title, headers, rows, widths, first=False, note=None):
    ws = wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = Font(name=FONT, italic=True, color='52667C', size=10)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.row_dimensions[1].height = 34
        ws.cell(row=1, column=1).alignment = Alignment(wrap_text=True, vertical='center', horizontal='right')
        r0 = 3
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r0, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=INK)
        cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    for i, row in enumerate(rows, r0 + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
            cell.border = Border(bottom=thin)
            if isinstance(v, str) and v.startswith('http'):
                cell.hyperlink = v; cell.font = Font(name=FONT, size=10, color=BLUE, underline='single')
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    if rows:
        ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
        t = Table(displayName=''.join(ch for ch in title if ch.isascii() and ch.isalnum()) or f"T{len(wb.worksheets)}", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(t)
    return ws

# ------------------------------------------------------------------ 1. section map
S = 'https://'
sections = [
 # page, section, source name, url, what we took, adaptation, motion, elementor
 ('كل الصفحات', 'الهيدر (يتحول لكبسولة)', 'tentwenty.com + طلب Ahmad (طريقة إبداعية للظهور والاختفاء)', S+'tentwenty.com/', 'الشكل العام (من شيت Website Inspiration): هيدر نظيف ومساحات واسعة', 'بدل ما يختفي عند النزول، الشريط يتجمع لكبسولة كحلية صغيرة في المنتصف (الشعار + زر «القائمة» + زر الحجز)، ويتفرد تاني عند الصعود أو المرور عليه أو الضغط على «القائمة»', 'تغير العرض والألوان بتوقيت expo-out', 'Header template + Sticky + CSS .is-compact + main.js'),
 ('كل الصفحات', 'الميجا منيو الإبداعي — الحلول', 'فكرة Ahmad + ltsthink.com + grigoletti.ch/en/projects', S+'ltsthink.com/', 'ltsthink: ميجا منيو ERP بثلاثة أعمدة معنونة + رابط "الكل". grigoletti projects: المعاينة تتغير مع المرور', 'ثلاث مجموعات + عمود معاينة حي + كارت دليل الفاتورة الإلكترونية. اللوحة تنفتح كدائرة من الزر الذي فتحها وتنطوي إليه، والصفحة الحالية لها خلفية زرقاء (طلب Ahmad)', 'clip-path circle من الزر + تبديل المعاينة', 'Elementor Mega Menu + Loop template لكل خدمة + HTML widget للمعاينات'),
 ('كل الصفحات', 'الميجا منيو — القطاعات', 'تصميم Innomate', '', 'شبكة مربعات للقطاعات', 'أيقونات تُرسم عند المرور + لوحة جانبية بوحدات القطاع وشعارات عملائه + مربع "كل القطاعات"', 'رسم خطوط الأيقونة', 'Mega Menu + Icon Box + JS صغير'),
 ('كل الصفحات', 'قائمة الموبايل (شاشة كاملة)', 'baunfire.com', S+'www.baunfire.com/', 'المنيو الجانبي/الشاشة الكاملة: روابط كبيرة وبيانات تواصل (من شيت الإلهام)', 'تنفتح كدائرة من زر القائمة، مجموعات قابلة للطي (الحلول/القطاعات/الشركة)', 'clip-path circle', 'Off-canvas / Popup'),
 ('كل الصفحات', 'منيو سفلي للموبايل (Dock)', 'a-lign.studio', S+'www.a-lign.studio/', 'المنيو السفلي (من شيت Website Inspiration) — لم يُفحص بصرياً، من الملاحظة وكود الصفحة', 'أربعة اختصارات: الحلول، القطاعات، احجز (كهرماني)، واتساب — يظهر بعد أول شاشة', 'انزلاق من الأسفل', 'Sticky container للموبايل فقط'),
 ('الرئيسية', 'الهيرو', 'ltsthink.com/erp-crm + neue.world + 800website.ae', S+'ltsthink.com/erp-crm/', 'ltsthink: لوحة تحكم HTML. neue.world: صف الثقة تحت العنوان', 'كبسولات الشعار خلف شاشة ERP حية؛ الكلمة المتغيرة تنزلق لأعلى داخل صندوق بعرض أطول كلمة، فالعنوان لا يتحرك أبداً (طلب Ahmad: كان بينط)', 'GSAP intro + word roll + parallax بالماوس', 'Container + HTML widget للشاشة + JS'),
 ('الرئيسية + العملاء', 'العملاء كصفوف (Projects)', 'grigoletti.ch/en/ (قسم Projects، طلب Ahmad)', S+'grigoletti.ch/en/', 'قسناه: صفوف بسنة وتصنيف بخط mono واسم ضخم وسهم؛ عند المرور الصف يغمق ويظهر thumbnail وشارة', 'بدل شريط الشعارات في الرئيسية (بعد القطاعات، شركة من كل قطاع) وبدل حائط الشعارات في صفحة العملاء (عمودين + فلتر). عند المرور: الصف كحلي، شعار العميل ينزلق، شارة «حلول القطاع»، والسهم كهرماني ويميل. الصف يفتح صفحة القطاع', 'CSS hover: width/scale/translate', 'Loop Grid بقالب صف + Custom CSS'),
 ('الرئيسية', 'لماذا نظام واحد (اقتباس يتراصّ)', 'a-lign.studio', S+'www.a-lign.studio/', 'طلب Ahmad بالظبط: قسم «People deserve better…». قسناه: قسم بارتفاع الشاشة، كلمات مبعثرة داخل أسطرها ثم تنزلق كلمة كلمة إلى كتلة مضبوطة الحواف، أقواس من شرائط، وسطر برتقالي يُكتب بحروف عشوائية', 'نفس البنية بنصنا (عربي/إنجليزي)؛ الكلمات تتحول من رمادي إلى كحلي عند استقرارها حتى تُقرأ، وعلامات الاقتباس بالكهرماني', 'FLIP + GSAP stagger عند الوصول لأعلى الشاشة + كتابة بحروف عشوائية', 'HTML widget + home.js'),
 ('الرئيسية + صفحة الحلول', 'الخدمات: الكروت المكدسة', 'grigoletti.ch', S+'grigoletti.ch/en/', 'فكرة Ahmad الثانية. قسناها في المتصفح: صفوف sticky بفرق 100px', 'العنوان والوصف ثابتان أعلى القسم (طلب Ahmad) والكروت تتكدس تحتهما؛ ارتفاع الكارت محسوب من ارتفاع الشاشة فيظهر كاملاً حتى على شاشة 720px', 'position: sticky + scale مرتبط بالتمرير', 'Containers sticky + Custom CSS + main.js'),
 ('الرئيسية', 'تفاصيل كارت الخدمة', 'grigoletti.ch/en/services + metikulous.design', S+'grigoletti.ch/en/services', 'grigoletti: عنوان + سطر + قائمة فرعية. metikulous: عدّاد الفصول 1/4', 'عدّاد 01/06 بالكهرماني + روابط الخدمات الفرعية كصفوف بأيقونة وسهم + شاشة المنتج', 'hover slide', 'Icon List / Loop'),
 ('من نحن', 'الأرقام (كروت على صفين)', 'pmmd.net + طلب Ahmad', S+'pmmd.net/', 'العدّاد الضخم في مقدمة pmmd', 'اتشال العدّاد المثبت من الرئيسية (كان بياخد 3.3 شاشة سكرول) بعد مراجعة impeccable/taste؛ الأرقام بقت في «لماذا Innomate» في الرئيسية وفي كروت «من نحن»', 'count-up عند الظهور', 'Counter widget'),
 ('الرئيسية', 'عرض وكلاء الذكاء الاصطناعي', 'surinder.design + بحث المنافسين', S+'surinder.design/', 'surinder: شريط "اسأل أي شيء" مثبت أسفل الواجهة', 'اختر وكيلاً (المخزون/التحصيل/فواتير الموردين/الإدارة) → السؤال يُكتب بالعامية → الوكيل يرد بالمصدر → إجراء ينتظر الموافقة. لا يوجد شريك ERPNext مصري لديه صفحة AI', 'typing + bubbles', 'HTML widget + home.js'),
 ('صفحة الفاتورة الإلكترونية', 'الفاتورة الإلكترونية ETA', 'ltsthink.com/erp-crm + بحث المنافسين', S+'ltsthink.com/erp-crm/', 'إشعار الفاتورة المعتمدة داخل واجهة النظام', 'اتشال من الرئيسية لأنه كان بيكرر كارت الخدمة 02؛ تفاصيله في صفحة الخدمة (المقارنة وأسباب الرفض)', '—', 'HTML widget'),
 ('الرئيسية', 'منهجية التطبيق (رحلة أفقية مثبتة)', 'pmmd.net/overview', S+'pmmd.net/overview', 'Journey أفقي مثبت (من شيت الإلهام)، والانتقال لقسم داكن كما في metikulous', 'خمس مراحل من رسم دورة العمل حتى الدعم، رسومات تُرسم عند الوصول، شريط تقدم', 'GSAP ScrollTrigger pin + horizontal scrub', 'Container + GSAP'),
 ('الرئيسية', 'القطاعات (قائمة مثبتة + مجسم 3D حي)', 'grigoletti.ch/en/projects + طلب Ahmad', S+'grigoletti.ch/en/projects', 'قائمة كبيرة بمعاينة تتغير', 'القسم يثبت؛ القائمة عمودين عشان تكفي شاشة لابتوب؛ مجسم القطاع حي (three.js): يتمايل ويطفو ويميل مع الماوس، وعند التغيير يلف ويصغر والجديد يقفز. كارت القطاع تحت المجسم', 'ScrollTrigger pin + live3d.js', 'Loop + JS؛ الصور WebP كبديل'),
 ('الرئيسية', 'التكاملات (مشهد 3D حي)', 'تصميم Innomate (طلب Ahmad: سكشن 3D)', '', 'شعار Innomate مجسّم في المركز', 'بلاطات التكاملات تدور على حلقتين مائلتين وحِزم بيانات تتجه للمركز؛ المشهد يميل مع المؤشر، والمرور على بلاطة يظهر اسمها؛ يتوقف خارج الشاشة', 'three.js (WebGL) + مؤشر', 'HTML widget + orbit3d.js (ES module)'),
 ('الرئيسية', 'لماذا Innomate + الأرقام', 'origin.global + ram.sa', S+'origin.global/', 'origin: "What makes us different". ram.sa: عنوان + شبكة', 'العنوان وبجانبه ثلاثة أرقام (50+ شركة، 7+ سنين، 20+ تكامل) فوق بينتو غير متساوٍ: محلي، مفتوح المصدر، 4–12 أسبوع، 24/7، ما بعد ERP. الكارت يرتفع بنعومة عند المرور بلا تغيير لون', 'count-up + lift', 'Grid container + Counter widget'),
 ('الرئيسية + العملاء', 'آراء العملاء', 'innomate-tech.com (الحالي)', S+'innomate-tech.com/', 'الشهادات الثلاث كما هي', 'الأولى بارزة داكنة + شعار الشركة (لاحقاً: عنوان بالنتيجة كما في surinder.design)', 'reveal', 'Testimonial widget'),
 ('كل الصفحات', 'الأسئلة الشائعة', 'grigoletti.ch + ram.sa', S+'grigoletti.ch/en/', 'أسئلة بعلامة + (grigoletti) وصياغة للبحث (ram.sa)', 'عمود جانبي sticky + زر واتساب + FAQPage schema', 'accordion', 'Accordion widget + schema'),
 ('كل الصفحات', 'الفوتر', 'origin.global', S+'origin.global/', 'origin: الفوتر كخريطة موقع كاملة', 'أعمدة الحلول والقطاعات والشركة (الشعار الضخم اتشال بطلب Ahmad)', '—', 'Footer template'),
 ('من نحن', 'الفريق', 'loop-agency.framer.website', S+'loop-agency.framer.website/', 'شبكة صور متدرجة الإزاحة (من شيت الإلهام)', 'ثلاثة أعمدة، العمود الأوسط منزاح لأسفل، ترتفع عند المرور — تنتظر الصور الحقيقية', 'hover lift', 'Loop Grid'),
 ('منهجية التطبيق', 'مخرجات كل مرحلة', 'a-lign.studio', S+'www.a-lign.studio/', 'مراحل بمخرجات واضحة', 'خط زمني عمودي، لكل مرحلة شارة مدة ومستندات تسليم', 'reveal', 'Timeline'),
 ('عملاؤنا', 'جدار العملاء مع فلتر القطاع', 'innomate-tech.com (الحالي)', S+'innomate-tech.com/', 'فكرة "العملاء حسب القطاع" الموجودة حالياً', 'شبكة 4×4 وأزرار فلترة', 'filter', 'Loop Grid + taxonomy filter'),
 ('تواصل معنا', 'النموذج بشرائح الاختيار', 'makeitwow.io', S+'makeitwow.io/', 'اختيار الاهتمامات كشرائح بدلاً من قائمة منسدلة', 'شرائح للخدمات العشر + تعبئة تلقائية من مقدّر الباقات', 'تحقق فوري', 'Form widget'),
 ('الباقات', 'الباقات + مقدّر المدة', 'LavaLoon + Plementus (من بحث المنافسين)', S+'plementus.com/', 'صفحة نموذج تسعير (LavaLoon) وأدوات تفاعلية (Plementus)', 'ثلاث باقات + مقدّر يعطي المدة والباقة (بدون أسعار حتى تُعتمد)', 'تحديث لحظي', 'Form + JS'),
 ('صفحات الخدمات', 'قالب صفحة الخدمة', 'ltsthink.com/erp-crm + grigoletti.ch/en/services + توصيات GEO', S+'grigoletti.ch/en/services', 'صفحة ERP: شاشة + خطوات + FAQ. grigoletti: كروت الخدمات', 'هيرو بشاشة المنتج + «باختصار» + ما يشمله التطبيق بجانب مجسم 3D للخدمة + الخطوات ككروت مكدسة مائلة تتفرد أفقياً مع التمرير + قبل/بعد + رأي العميل في شريط مستقل + التكلفة + FAQ + خدمات مرتبطة (نفس الكروت)', 'fan scrub + live 3D', 'Single template + ACF fields'),
 ('صفحات القطاعات', 'قالب صفحة القطاع', 'توصيات بحث المنافسين', '', 'مصطلحات القطاع + جدول المراحل', 'هيرو بمجسم القطاع 3D حي + شعارات عملاء القطاع + المشاكل ككروت مكدسة تتفرد + جدول الوحدات', 'live 3D + fan', 'Single template'),
 ('صفحات الحلول والقطاعات', 'هيرو صفحات الفهرس', 'تصميم Innomate', '', '—', 'أيقونات الخدمات/القطاعات مرصوصة في ثلاثة أعمدة بشكل كبسولات الشعار', 'ظهور متتابع', 'Icon widgets'),
 ('كل الصفحات', 'التمرير الناعم وتوقيت الحركة', 'a-lign.studio', S+'www.a-lign.studio/', 'a-lign يستخدم Lenis للتمرير الناعم', 'Lenis يقود GSAP ScrollTrigger، فكل الحركات المرتبطة بالتمرير على نفس الخط الزمني، وتوقيت expo-out أبطأ وأنعم', 'Lenis + ScrollTrigger', 'سكربت واحد في القالب'),
 ('الرئيسية + العملاء', 'آراء العملاء (تبويبات + اقتباس واحد)', 'طلب Ahmad (الكروت كانت متداخلة)', '', '—', 'اقتباس واحد كبير على المسرح؛ العملاء الثلاثة تبويبات بشعار واسم، وخط تقدم كهرماني يمر ثم ينتقل للاقتباس التالي؛ المرور يوقفه، والاقتباس يربط بصفحة الخدمة', 'crossfade + CSS animation للخط', 'Tabs widget أو Loop Carousel + Custom CSS'),
 ('كيف نعمل', 'المراحل (رحلة أفقية مثبتة)', 'pmmd.net + طلب Ahmad', S+'pmmd.net/', 'نفس رحلة الرئيسية', 'المراحل الخمس بنفس رسومات الرئيسية + ما تستلمه في نهاية كل مرحلة؛ القسم يثبت والمسار يتحرك أفقياً', 'ScrollTrigger pin + translateX', 'Horizontal scroll container + main.js'),
 ('كل الصفحات', 'قسم الحجز (CTA)', 'تصميم Innomate + طلب Ahmad', '', '—', 'كارت كحلي؛ عند المرور ينتشر ضوء أزرق طبيعي من نقطة بيضاء في منتصف الحافة السفلية: قبة ناعمة قوية في النص وخافتة عند الأطراف، بلا أي خط', 'radial gradients + blur + screen blend', 'Container + Custom CSS (::before/::after)'),
 ('كل الصفحات الداخلية', 'كروت مكدسة مائلة تتفرد أفقياً', 'grigoletti.ch/en/services (طلب Ahmad)', S+'grigoletti.ch/en/services', 'قسناها: ثلاث كروت فوق بعض بميل خفيف تتحرك لأماكنها في صف مع التمرير؛ عنوان كبير، وصف، سطور بفواصل (Timeline / Learn more)، ونقاط موضع أعلى الكارت', 'خطوات الخدمة، الباقات، الخدمات المرتبطة، مشاكل القطاع، ما تحتاجه منك (كيف نعمل). النقاط تبين ترتيب الكارت', 'ScrollTrigger scrub: x/y/rotation من المنتصف', 'Container + main.js (data-fan)'),
 ('كل الصفحات', 'كروت بهوفر رأسي وتغيير لون', 'طلب Ahmad (أرقام صفحة من نحن)', '', '—', 'أرقام «من نحن» على صفين داخل كروت؛ عند المرور اللون الكحلي يطلع من تحت لفوق، الكارت يرتفع، والرقم يلف لأعلى ويظهر بالكهرماني. نفس الهوفر على كروت القيم والخدمات والقطاعات والتكاملات ولماذا Innomate', 'scaleY من الحافة السفلية + text roll', 'Custom CSS (.hcard)'),
 ('كل الصفحات الداخلية', 'رسومات 3D للصفحات', 'تصميم Innomate (طلب Ahmad: الصفحات مكدسة ومحتاجة Illustration)', '', '—', 'مجسم لكل خدمة (10) ولكل صفحة (من نحن، الأسعار، التكاملات، العملاء، الوظائف، المقالات، الأسئلة، …) من نفس ملف النماذج؛ حي في الصفحة وصورة WebP تحته', 'live3d.js: تمايل + طفو + ميل مع الماوس', 'Image widget بالـ WebP، أو HTML widget للنسخة الحية'),
 ('كل الصفحات', 'زر واتساب عائم', 'طلب Ahmad', '', '—', 'دائرة واتساب بنبض خفيف تظهر بعد أول شاشة، وتتمدد عند المرور لتظهر «كلّمنا على واتساب». على الموبايل الواتساب في الـ Dock', 'spring in + ring pulse', 'Floating button / Sticky container'),
 ('كل الصفحات', 'ظهور الأقسام', 'طلب Ahmad (دخول بسيط وأنيق)', '', '—', 'العناوين تطلع من خط (قناع من أسفل)، الكروت تصعد بتتابع', 'clip-path + translate، expo-out', 'Motion Effects: Entrance'),
]
ws = sheet('خريطة السكاشن', ['الصفحة', 'السكشن', 'المصدر', 'الرابط', 'ماذا أخذنا', 'كيف طوّعناه لـ Innomate', 'الحركة', 'التنفيذ في Elementor'], sections,
           [16, 30, 30, 34, 40, 60, 26, 34], first=False,
           note='كل سكشن في النموذج: من أين جاءت الفكرة، وماذا أخذنا منها (نمط وتفاعل فقط — لا كود ولا صور منسوخة)، وكيف تغيّرت لتناسب Innomate.')

# ------------------------------------------------------------------ 2. reference review
if refs:
    rev_rows = [(r.get('site', ''), r.get('url', ''), 'نعم' if r.get('reachable') in (True, 'yes', 'Yes') else ('لا' if r.get('reachable') in (False, 'no', 'No') else str(r.get('reachable', ''))),
                 r.get('theme', ''), r.get('borrowed_section', ''), r.get('innomate_section', ''), r.get('what_we_take', ''), r.get('adaptation', ''), r.get('motion', '')) for r in refs]
else:
    rev_rows = []
sheet('مراجعة المواقع', ['الموقع', 'الرابط', 'متاح؟', 'الثيم', 'السكشن المرشح', 'مكانه عندنا', 'ما يستحق الأخذ', 'التطويع', 'الحركة'], rev_rows,
      [22, 32, 8, 10, 30, 26, 50, 50, 34], note='مراجعة كل موقع في nitaq-reference-sites.xlsx وWebsite Inspiration.xlsx (المصدر: research/references.md).')

# ------------------------------------------------------------------ 3. competitors
comp = [
 ('Data Value Solutions', 'https://datavalue.solutions', 'ERPNext (شريك فضي)', 'مصر + السعودية', 'نعم (4 + وحدات)', '6', 'متوسط', 'لا', 'أزرق فاتح', 'أقوى منافس مباشر؛ صفحة ETA قصيرة'),
 ('Operion (Smart Solutions)', 'https://operionss.com', 'ERPNext', 'مصر + السعودية', 'نعم (9)', '9', 'ذكر فقط', 'دليل بلا أرقام', 'بنفسجي فاتح', 'الأكمل SEO؛ مقارنات مع Odoo/SAP'),
 ('LavaLoon', 'https://lavaloon.com', 'ERPNext (شريك فضي)', 'مصر + السعودية', 'نعم', '6', 'السعودية فقط', 'نعم (نموذج)', 'فاتح، أحمر/كهرماني', 'أفضل نموذج لصفحة تسعير'),
 ('Go-Live Solutions', 'https://golive-solutions.com', 'ERPNext', 'مصر', 'جزئياً', 'لا', 'صفحة منتج', 'لا', 'فاتح، صور ستوك', 'أقرب مزيج خدمات لـ Innomate، موقع قديم'),
 ('Plementus', 'https://plementus.com', 'Odoo (ذهبي)', 'مصر/الإمارات/السعودية', 'نعم (6)', '12', 'FAQ', 'حاسبة ROI', 'فاتح، برقوقي/حبر', 'الأفضل بناءً: اختبار جاهزية، FAQ schema، llms.txt'),
 ('Daftra', 'https://daftra.com', 'SaaS', 'مصر والخليج', 'نعم', '50+', 'عالٍ', 'نعم', 'أزرق فاتح', 'فواتير متحركة في العرض'),
 ('Wafeq', 'https://wafeq.com/ar-eg', 'SaaS', 'مصر والخليج', 'نعم', '—', 'عالٍ (دليل)', 'نعم', 'نيلي فاتح', 'كتّاب بأسماء (E-E-A-T)'),
 ('Pioneers Solutions', 'https://pioneers-solutions.com', 'محلي', 'مصر', 'نعم', 'بعض', 'الأعمق (مقالات)', 'لا', 'فاتح، قديم', 'أقوى محتوى عربي عن الفاتورة الإلكترونية'),
 ('Innomate (الحالي)', 'https://innomate-tech.com', 'ERPNext', 'مصر', 'لا (صفحة واحدة)', 'لا', 'لا يوجد', 'لا', 'كحلي داكن', 'المحتوى يُولَّد بالجافاسكربت فلا تراه محركات البحث'),
]
sheet('المنافسون', ['المنافس', 'الموقع', 'المنصة', 'السوق', 'صفحات خدمات منفصلة', 'صفحات قطاعات', 'عمق محتوى ETA', 'أسعار معلنة', 'الثيم', 'ملاحظة'], comp,
      [24, 30, 18, 18, 18, 12, 16, 14, 18, 44], note='ملخص من research/competitors.md (14 منافساً + 3 مراجع عالمية). الأرقام كما ظهرت على مواقعهم في 30 سبتمبر 2026.')

# ------------------------------------------------------------------ 4. keywords
kw_rows = [('الرئيسية', 'نظام ERP مصر · ERPNext مصر', 'شريك ERPNext، برنامج ERP للشركات، سيستم ERP', 'ERPNext partner Egypt', 'تجاري', 'index.html')]
for s in content['services']:
    kw_rows.append((s['name'], s['kw'][0] + ' · ' + s['kw'][1], '، '.join(s['kw'][2:]), ' · '.join(s.get('enKw', [])[:3]), 'تجاري/معلوماتي', f"services/{s['slug']}.html  ·  en/services/{s['slug']}.html"))
for s in content['sectors']:
    kw_rows.append((s['name'], s['title'].split('|')[0].strip(), '', '', 'تجاري', f"sectors/{s['slug']}.html"))
kw_rows += [
 ('الباقات', 'سعر ERPNext في مصر · تكلفة تطبيق نظام ERP', 'ERPNext بكام، هل ERPNext مجاني', 'ERPNext pricing Egypt', 'تجاري', 'pricing.html'),
 ('التكاملات', 'ربط ERPNext مع الأنظمة', 'ربط الفاتورة الإلكترونية، ربط جهاز البصمة، ربط Power BI', 'ERPNext integrations', 'تجاري', 'integrations.html'),
 ('الأسئلة الشائعة', 'أسئلة شائعة عن ERPNext', 'هل ERPNext مجاني، هل ERPNext يدعم العربي', 'ERPNext FAQ', 'معلوماتي', 'faq.html'),
 ('الوظائف', 'وظائف ERPNext مصر', 'وظائف العاشر من رمضان IT', 'ERPNext jobs Egypt', 'تنقلي', 'careers.html'),
]
sheet('الكلمات المفتاحية', ['الصفحة', 'الكلمة الأساسية', 'كلمات ثانوية', 'بالإنجليزية', 'نية البحث', 'الرابط'], kw_rows, [30, 44, 50, 26, 16, 34],
      note='أحجام البحث تقديرية (Ahrefs غير متصل في هذه الجلسة) — راجعها في Keyword Planner قبل الاعتماد. في مصر "برنامج" أكثر بحثاً من "نظام" للشركات الصغيرة.')

# ------------------------------------------------------------------ 5. pages
pg_rows = []
for p in sorted(content['pages'], key=lambda x: (x['path'].count('/'), x['path'])):
    pg_rows.append((p['path'], p['h1'], p['title'], len(p['title']), p['desc'], len(p['desc']), p['schema']))
sheet('صفحات الموقع', ['الملف', 'H1', 'Title', 'طول العنوان', 'Meta description', 'طول الوصف', 'Schema'], pg_rows, [34, 44, 60, 10, 80, 10, 40],
      note='كل صفحة (عربي في الجذر، إنجليزي تحت /en/) لها عنوان ووصف وH1 واحد فريد، وhreflang يربط كل صفحة بنظيرتها. الطول المثالي: العنوان ≤ 70 حرفاً، الوصف 120–160.')

# ------------------------------------------------------------------ 6. placeholders
ph = [
 ('صور وأسماء الفريق', 'about.html', 'صور بخلفية فاتحة موحدة + الاسم والمسمى', 'Innomate'),
 ('دراسات الحالة', 'clients.html', 'Electro Metal، Trend Pharm، Naturetex — تُكتب بموافقة العميل ورقم نتيجة حقيقي', 'Innomate + العملاء'),
 ('أسعار الباقات بالجنيه', 'pricing.html', '"يبدأ من" لكل باقة', 'Innomate'),
 ('روابط LinkedIn وFacebook', 'كل الصفحات (الفوتر + Schema)', 'تأكيد الحسابات الرسمية', 'Innomate'),
 ('مراجعة النسخة الإنجليزية', '/en/ (37 صفحة)', 'قراءة لغوية من متحدث أصلي، وتأكيد ترجمة شهادات العملاء الثلاث مع أصحابها', 'Innomate'),
 ('نموذج التواصل', 'contact.html', 'ربطه بـ CRM أو بريد (حالياً تحقق في المتصفح فقط)', 'عند التحويل لـ WordPress'),
 ('تواريخ مراحل الفاتورة/الإيصال الإلكتروني', 'services/eta-e-invoice.html', 'لم نضع جدولاً زمنياً عمداً — يُضاف بعد مراجعته على eta.gov.eg', 'Innomate'),
 ('بيانات الشاشات التوضيحية', 'كل الـ mocks', 'أرقام وأسماء توضيحية وليست بيانات عملاء', 'للعلم'),
 ('الموديلات المدعومة من ZKTeco', 'services/zkteco-attendance.html', 'قائمة الأجهزة المجربة فعلاً', 'Innomate'),
 ('حالة شراكة Frappe', 'about.html + llms.txt', 'إن وجدت شراكة رسمية تُذكر بدقة', 'Innomate'),
 ('المقالات', 'resources.html', '8 عناوين مخططة، أولها دليل الفاتورة الإلكترونية', 'فريق المحتوى'),
 ('صور القطاعات', 'الرئيسية (القطاعات)', 'مجسمات 3D توضيحية؛ مراجعة AI slop تنصح بصور حقيقية من مصانع العملاء عند توفرها', 'Innomate'),
 ('الفريق ودراسات الحالة', 'about.html / clients.html', 'مخفية في النسخة الحالية حتى يتوفر محتوى حقيقي (DRAFTS=1 npm run build لإظهارها)', 'Innomate'),
]
sheet('مؤقت للتأكيد', ['العنصر', 'أين', 'المطلوب', 'المسؤول'], ph, [34, 34, 70, 22],
      note='محتوى مؤقت لا يُنشر كما هو.')

# ------------------------------------------------------------------ summary (first sheet)
sm = wb.create_sheet('ملخص', 0)
sm.sheet_view.rightToLeft = True
lines = [
 ('Innomate — مصادر التصميم والمحتوى', None),
 ('تاريخ الإعداد: 30 سبتمبر 2026', None),
 ('', None),
 ('عدد السكاشن الموثقة', "=COUNTA('خريطة السكاشن'!B4:B200)"),
 ('سكاشن مأخوذة من مواقع مرجعية (لها رابط)', "=COUNTIF('خريطة السكاشن'!D4:D200,\"http*\")"),
 ('المواقع المراجَعة', "=COUNTA('مراجعة المواقع'!A4:A200)"),
 ('صفحات الموقع', "=COUNTA('صفحات الموقع'!A4:A200)"),
 ('عناوين أطول من 65 حرفاً', "=COUNTIF('صفحات الموقع'!D4:D200,\">65\")"),
 ('عناصر مؤقتة تنتظر التأكيد', "=COUNTA('مؤقت للتأكيد'!A4:A200)"),
 ('', None),
 ('القرارات الأساسية', None),
 ('• ثيم نهاري: خلفية #F4F7FA، نص كحلي #0A1A2B، أزرق الشعار #388CCC، والكهرماني #F3A93A لون الأزرار (كل المنافسين أزرق).', None),
 ('• الخطوط: Readex Pro للعناوين، IBM Plex Sans Arabic للنص، IBM Plex Mono للأرقام.', None),
 ('• الموتيف: كبسولات الشعار الثلاث في كل مكان (الهيرو، الميجا منيو، الأزرار، اللودر، الـ CTA).', None),
 ('• فكرتا Ahmad: الميجا منيو الإبداعي، والكروت المكدسة (grigoletti.ch).', None),
]
for i, (a, b) in enumerate(lines, 1):
    c = sm.cell(row=i, column=1, value=a); c.font = Font(name=FONT, size=11, bold=(i in (1, 11)), color=INK if i != 1 else BLUE)
    c.alignment = Alignment(horizontal='right', wrap_text=True)
    if b:
        v = sm.cell(row=i, column=2, value=b); v.font = Font(name=FONT, size=11, bold=True); v.alignment = Alignment(horizontal='center')
sm['A1'].font = Font(name=FONT, size=16, bold=True, color=BLUE)
sm.column_dimensions['A'].width = 95; sm.column_dimensions['B'].width = 14

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
out = os.path.join(HERE, 'innomate-section-sources.xlsx')
wb.save(out)
print('saved', out, 'refs:', len(refs))
